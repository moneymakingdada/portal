"""
The downloadable sample letter of authorization (a .docx the customer edits,
signs and uploads with their sender ID request). Built on demand so it can be
pre-filled with what we already know about the signed-in user; anything we
don't know is left as a highlighted [placeholder] for them to complete.
"""
import io
import re

from django.conf import settings
from django.utils import timezone
from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.shared import Pt

ADDRESSEE = ("The Products Manager,", "AirtelTigo Ghana,", "Barnes Rd, Accra, Ghana")
CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

_PLACEHOLDER = re.compile(r"(\[[^\]]+\])")


def _add(doc, text: str = "", *, bold: bool = False, space_after: int = 8):
    """Add a paragraph; any [bracketed] text is highlighted so it's obvious it needs editing."""
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(space_after)
    for chunk in _PLACEHOLDER.split(text):
        if not chunk:
            continue
        run = paragraph.add_run(chunk)
        run.bold = bold
        if _PLACEHOLDER.fullmatch(chunk):
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW
    return paragraph


def build_letter(*, business: bool, sender_id: str = "", full_name: str = "", email: str = "",
                 phone: str = "", organization: str = "") -> bytes:
    partner = settings.SENDER_LETTER_PARTNER
    platform = settings.APP_NAME
    name = full_name or "[Your full name]"
    sender = sender_id or "[Your sender ID]"
    company = organization or "[Company name]"

    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)

    if business:
        _add(doc, "[Print this on your company letterhead]")
    today = timezone.localdate()
    _add(doc, f"{today.day} {today:%B %Y}", space_after=14)
    for index, line in enumerate(ADDRESSEE):
        _add(doc, line, space_after=0 if index < len(ADDRESSEE) - 1 else 14)
    _add(doc, "Dear Sir/Madam,", space_after=10)

    if business:
        _add(doc, f"I, {name}, in my capacity as [Your position, e.g. Managing Director] of {company}, confirm that "
                  f"{company} has the authority to use the following sender ID(s) through the {platform} SMS "
                  "platform:")
    else:
        _add(doc, f"I, {name}, in my capacity as [your role, e.g. Owner / Freelance Developer], confirm that I have "
                  f"the authority to use the following sender ID(s) through the {platform} SMS platform:")
    _add(doc, f"Sender ID: {sender}", bold=True)

    who = "us" if business else "me"
    _add(doc, f"I write this letter to request that the above sender ID(s) be whitelisted for {platform}, who serves "
              f"as {'our' if business else 'my'} authorized partner for sending SMS messages to "
              f"{'our' if business else 'my'} contacts.")
    _add(doc, "This sender ID will be used for sending [describe what you will send, e.g. OTPs, transaction alerts "
              "and customer notifications].")
    if business:
        _add(doc, f"A copy of {company}'s business registration certificate is attached.")
    else:
        _add(doc, "A copy of my Ghana Card is attached.")

    _add(doc, "Below is my contact information:")
    _add(doc, f"Name: {name}", space_after=0)
    _add(doc, "Position: [Your position]", space_after=0)
    if business:
        _add(doc, f"Company: {company}", space_after=0)
    _add(doc, f"Email: {email or '[Your email address]'}", space_after=0)
    _add(doc, f"Phone: {phone or '[Your phone number]'}", space_after=14)

    _add(doc, f"Please contact {who} if you need more information or clarification.")
    _add(doc, "Thank you for your attention to this matter.", space_after=22)
    _add(doc, "Sincerely,", space_after=26)
    _add(doc, "______________________________", space_after=0)
    _add(doc, name)
    if business:
        _add(doc, "[Company stamp, if you have one]")

    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()
