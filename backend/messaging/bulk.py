"""
Reading a list of recipients for a bulk send.

Three ways in, all producing the same thing:
  * numbers pasted into a box, separated by commas, semicolons, colons, pipes
    or new lines (spaces *inside* a number are fine: "024 123 4567");
  * a CSV or plain-text file;
  * an Excel (.xlsx) file, including Arkesel's contacts-import layout
    (Phone Number, Email, Company, First Name, Last Name, Birth Date).

Nothing here touches the database or sends anything: it only turns input into
a clean, de-duplicated list of recipients plus a list of what was rejected and why.
"""
import csv
import datetime
import io
import re
import zipfile
from dataclasses import dataclass, field

from django.core.exceptions import ValidationError
from django.core.validators import validate_email

from common.errors import ApiError
from common.phone import InvalidPhone, normalize_gh_number

MAX_RECIPIENTS = 2000            # per bulk send
MAX_FILE_BYTES = 2 * 1024 * 1024
MAX_PASTE_CHARS = 200_000
MAX_SCANNED = 20_000             # rows/entries read before giving up on a runaway file
MAX_UNZIPPED_BYTES = 50 * 1024 * 1024
MAX_REPORTED_INVALID = 25

_SEPARATORS = re.compile(r"[,;:|\r\n\t]+")
_SPACES = re.compile(r"\s+")

_DATE_FORMATS = ("%d-%m-%Y", "%d/%m/%Y", "%d.%m.%Y", "%Y-%m-%d", "%Y/%m/%d")


@dataclass
class Entry:
    """One line or row as found, before the number is checked."""
    value: str
    where: str
    name: str = ""
    email: str = ""
    birthday: datetime.date | None = None


@dataclass
class Parsed:
    recipients: list = field(default_factory=list)   # [{"phone", "name", "email", "birthday"(ISO str or "")}]
    invalid: list = field(default_factory=list)      # [{"value", "reason", "where"}] (capped)
    invalid_count: int = 0
    duplicates: int = 0
    found: int = 0
    warnings: dict = field(default_factory=dict)     # e.g. {"bad_birth_dates": 2}


# ---------------------------------------------------------------------------
# Pasted text
# ---------------------------------------------------------------------------
def entries_from_text(text: str) -> list[Entry]:
    entries = []
    for line_no, line in enumerate((text or "").splitlines(), start=1):
        for chunk in _SEPARATORS.split(line):
            chunk = chunk.strip()
            if chunk:
                entries.append(Entry(value=chunk, where=f"line {line_no}"))
        if len(entries) > MAX_SCANNED:
            raise _too_many_rows()
    return entries


# ---------------------------------------------------------------------------
# Files
# ---------------------------------------------------------------------------
def entries_from_file(upload) -> tuple[list[Entry], dict]:
    """Read a .csv, .txt or .xlsx upload. Returns (entries, warnings)."""
    name = (getattr(upload, "name", "") or "").lower()
    if getattr(upload, "size", 0) > MAX_FILE_BYTES:
        raise ApiError("file_too_large", "That file is too large. Keep it under 2 MB.", 413)
    data = upload.read(MAX_FILE_BYTES + 1)
    if len(data) > MAX_FILE_BYTES:
        raise ApiError("file_too_large", "That file is too large. Keep it under 2 MB.", 413)
    if not data.strip():
        raise ApiError("empty_file", "That file is empty.")

    if name.endswith(".txt"):
        # A plain text file is just a longer version of the paste box, so it follows the same separator rules.
        return entries_from_text(_decode(data)), {}
    if name.endswith(".xlsx"):
        rows = _xlsx_rows(data)
    elif name.endswith(".csv"):
        rows = _csv_rows(data)
    elif name.endswith(".xls"):
        raise ApiError("unsupported_file", "Old .xls files aren't supported. In Excel choose Save As and pick .xlsx or .csv.")
    else:
        raise ApiError("unsupported_file", "Upload a .csv, .txt or .xlsx file.")
    return _entries_from_rows(rows)


def _xlsx_rows(data: bytes):
    if not data.startswith(b"PK"):
        raise ApiError("unreadable_file", "That doesn't look like a real .xlsx file.")
    try:
        # A tiny file can expand enormously; look at the declared sizes before opening it.
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            if sum(i.file_size for i in archive.infolist()) > MAX_UNZIPPED_BYTES:
                raise ApiError("unreadable_file", "That spreadsheet is too large to read.")
        from openpyxl import load_workbook

        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        sheet = workbook.worksheets[0]
        rows = []
        for row in sheet.iter_rows(values_only=True):
            rows.append(list(row))
            if len(rows) > MAX_SCANNED:
                raise _too_many_rows()
        workbook.close()
        return rows
    except ApiError:
        raise
    except Exception as exc:   # corrupt zip, missing sheet, bad xml...
        raise ApiError("unreadable_file", "We couldn't read that spreadsheet. Save it again as .xlsx or .csv and retry.") from exc


def _decode(data: bytes) -> str:
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        return data.decode("cp1252", errors="replace")


def _csv_rows(data: bytes):
    text = _decode(data)
    first_line = next((ln for ln in text.splitlines() if ln.strip()), "")
    # Excel in some regions saves "CSV" with semicolons; pick whichever separator the first line uses most.
    delimiter = max(",;\t|", key=lambda d: first_line.count(d)) if any(d in first_line for d in ",;\t|") else ","
    rows = []
    try:
        for row in csv.reader(io.StringIO(text), delimiter=delimiter):
            rows.append(row)
            if len(rows) > MAX_SCANNED:
                raise _too_many_rows()
    except csv.Error as exc:
        raise ApiError("unreadable_file", "We couldn't read that file. Check it's a plain CSV and retry.") from exc
    return rows


def _norm_header(cell) -> str:
    return re.sub(r"[^a-z0-9]", "", str(cell or "").lower())


def _classify_header(header: str) -> str | None:
    """Map a column heading to phone / email / birthday / first / last / name / None."""
    if not header or any(word in header for word in ("company", "business", "organi", "employer")):
        return None
    if "mail" in header:
        return "email"
    if "birth" in header or header in {"dob", "bday"}:
        return "birthday"
    if "name" in header or header in {"surname", "first", "last", "fullname"}:
        if header.startswith("first") or header == "forename" or header == "givenname":
            return "first"
        if header.startswith("last") or header == "surname" or header == "familyname":
            return "last"
        return "name"
    if any(word in header for word in ("phone", "mobile", "msisdn", "cell", "number", "contact", "tel")):
        return "phone"
    return None


def _cell_text(cell) -> str:
    """A cell as clean text. Excel stores numbers as numbers (and drops a leading zero), so render
    whole numbers without a trailing ".0" or exponent."""
    if cell is None:
        return ""
    if isinstance(cell, bool):
        return ""
    if isinstance(cell, int):
        return str(cell)
    if isinstance(cell, float):
        return str(int(cell)) if cell.is_integer() else str(cell)
    return str(cell).strip()


def _entries_from_rows(rows) -> tuple[list[Entry], dict]:
    # Keep each row's real position in the file so a problem can be reported as "row 7" (blank rows included).
    numbered = [(n, r) for n, r in enumerate(rows, start=1) if any(_cell_text(c) for c in r)]
    if not numbered:
        raise ApiError("empty_file", "That file has no rows in it.")

    columns: dict[str, int] = {}
    for index, cell in enumerate(numbered[0][1]):
        kind = _classify_header(_norm_header(cell))
        if kind and kind not in columns:
            columns[kind] = index
    has_header = "phone" in columns
    if has_header:
        body = numbered[1:]
    else:
        # No recognisable heading row: treat the first column as the numbers, header or not.
        columns, body = {"phone": 0}, numbered

    entries, bad_dates, bad_emails = [], 0, 0
    for offset, row in body:
        def cell(kind):
            index = columns.get(kind)
            return _cell_text(row[index]) if index is not None and index < len(row) else ""

        raw_phone = cell("phone")
        raw_birth = row[columns["birthday"]] if "birthday" in columns and columns["birthday"] < len(row) else None
        if not raw_phone and not any(cell(k) for k in ("name", "first", "last", "email")):
            continue
        name = cell("name") or " ".join(p for p in (cell("first"), cell("last")) if p)
        email = cell("email")
        if email:
            try:
                validate_email(email)
            except ValidationError:
                email, bad_emails = "", bad_emails + 1
        birthday = parse_birthday(raw_birth)
        if birthday is None and _cell_text(raw_birth):
            bad_dates += 1
        entries.append(Entry(value=raw_phone, where=f"row {offset}", name=name[:120], email=email, birthday=birthday))
        if len(entries) > MAX_SCANNED:
            raise _too_many_rows()

    warnings = {}
    if bad_dates:
        warnings["bad_birth_dates"] = bad_dates
    if bad_emails:
        warnings["bad_emails"] = bad_emails
    return entries, warnings


def parse_birthday(raw) -> datetime.date | None:
    """Accepts real dates (from Excel) and day-first text like 31-12-2000 or 31/12/2000, plus ISO dates."""
    if isinstance(raw, datetime.datetime):
        found = raw.date()
    elif isinstance(raw, datetime.date):
        found = raw
    else:
        text = _cell_text(raw)
        found = None
        for fmt in _DATE_FORMATS:
            try:
                found = datetime.datetime.strptime(text, fmt).date()
                break
            except ValueError:
                continue
    if found is None or found.year < 1900 or found > datetime.date.today():
        return None
    return found


# ---------------------------------------------------------------------------
# Checking numbers and removing duplicates
# ---------------------------------------------------------------------------
def _normalize_chunk(value: str) -> list[str]:
    """Return the normalized number(s) in `value`. A chunk like "0241234567 0551234567"
    (two numbers split by a space) is accepted when each half is a valid number; a
    chunk like "024 123 4567" is one number. Raises InvalidPhone when it is neither."""
    try:
        return [normalize_gh_number(value)]
    except InvalidPhone:
        parts = [p for p in _SPACES.split(value.strip()) if p]
        if len(parts) > 1:
            try:
                return [normalize_gh_number(p) for p in parts]
            except InvalidPhone:
                pass
        raise


def build(entries: list[Entry], warnings: dict | None = None) -> Parsed:
    parsed = Parsed(warnings=dict(warnings or {}))
    seen: dict[str, dict] = {}
    for entry in entries:
        parsed.found += 1
        if not entry.value:
            _reject(parsed, entry, "Missing phone number.")
            continue
        try:
            phones = _normalize_chunk(entry.value)
        except InvalidPhone:
            _reject(parsed, entry, "Not a valid Ghana mobile number.")
            continue
        for phone in phones:
            existing = seen.get(phone)
            if existing is not None:
                parsed.duplicates += 1
                # A later row can still add details the first one lacked.
                for key in ("name", "email"):
                    if not existing[key] and getattr(entry, key) and len(phones) == 1:
                        existing[key] = getattr(entry, key)
                if not existing["birthday"] and entry.birthday and len(phones) == 1:
                    existing["birthday"] = entry.birthday.isoformat()
                continue
            seen[phone] = {
                "phone": phone,
                "name": entry.name if len(phones) == 1 else "",
                "email": entry.email if len(phones) == 1 else "",
                "birthday": entry.birthday.isoformat() if entry.birthday and len(phones) == 1 else "",
            }
    parsed.recipients = list(seen.values())
    if len(parsed.recipients) > MAX_RECIPIENTS:
        raise ApiError(
            "too_many_recipients",
            f"That's {len(parsed.recipients):,} numbers. You can send to up to {MAX_RECIPIENTS:,} at a time. Split the list and send it in parts.",
        )
    return parsed


def _reject(parsed: Parsed, entry: Entry, reason: str):
    parsed.invalid_count += 1
    if len(parsed.invalid) < MAX_REPORTED_INVALID:
        parsed.invalid.append({"value": entry.value[:40], "reason": reason, "where": entry.where})


def _too_many_rows():
    return ApiError("too_many_recipients", "That list is far too long. You can send to up to 2,000 numbers at a time.")


def read_recipients(*, text: str = "", upload=None) -> Parsed:
    """Pasted numbers and/or an uploaded file, merged into one clean list."""
    if len(text or "") > MAX_PASTE_CHARS:
        raise ApiError("paste_too_long", "That's too much text to paste. Upload it as a file instead.")
    entries = entries_from_text(text)
    warnings: dict = {}
    if upload is not None:
        file_entries, warnings = entries_from_file(upload)
        entries = file_entries + entries     # the file's details (names, birthdays) win over a pasted duplicate
    if not entries:
        raise ApiError("no_recipients", "Paste some phone numbers or upload a file first.")
    return build(entries, warnings)


# ---------------------------------------------------------------------------
# Sample file
# ---------------------------------------------------------------------------
SAMPLE_CSV = (
    "Phone Number,Email,Company,First Name,Last Name,Birth Date\r\n"
    "0240000000,ama@example.com,Adom Bakery,Ama,Mensah,31-12-1990\r\n"
    "0551234567,,,Kofi,Boateng,\r\n"
    "0201112222,,,,,\r\n"
)
